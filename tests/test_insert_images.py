from pathlib import Path

import pytest
from docx import Document
from PIL import Image

from insert_images import (
    STATUS_DUPLICATE_NAME,
    STATUS_NO_IMAGE,
    STATUS_NO_TABLE,
    STATUS_NORMAL,
    process_document,
)


def make_image(path: Path) -> Path:
    Image.new("RGB", (60, 40), "red").save(path)
    return path


def make_docx(path: Path, names: list[str], extra_shapes: list[tuple[int, int]] = ()) -> Path:
    document = Document()
    for name in names:
        table = document.add_table(rows=2, cols=1)
        table.rows[0].cells[0].text = name
        table.rows[1].cells[0].text = "기존 내용"
    for rows, cols in extra_shapes:
        table = document.add_table(rows=rows, cols=cols)
        table.rows[0].cells[0].text = "a"
    document.save(str(path))
    return path


def count_pictures(path: Path) -> int:
    return len(Document(str(path)).inline_shapes)


def second_row_text(path: Path, table_index: int) -> str:
    return Document(str(path)).tables[table_index].rows[1].cells[0].text


def test_inserts_all_matching_images(tmp_path):
    images = [make_image(tmp_path / "a.png"), make_image(tmp_path / "B.jpg")]
    source = make_docx(tmp_path / "in.docx", ["a", " b "])
    output = tmp_path / "out.docx"

    report = process_document(source, images, output)

    assert report.ok and report.saved
    assert (report.image_count, report.target_count, report.normal_count, report.exception_count) == (2, 2, 2, 0)
    assert count_pictures(output) == 2
    assert "기존 내용" not in second_row_text(output, 0)


def test_image_count_ignores_non_image_files(tmp_path):
    images = [make_image(tmp_path / "a.png"), tmp_path / "note.txt"]
    source = make_docx(tmp_path / "in.docx", ["a"])

    report = process_document(source, images, tmp_path / "out.docx")

    assert report.image_count == 1


def test_table_without_image_is_exception_and_not_saved(tmp_path):
    images = [make_image(tmp_path / "a.png")]
    source = make_docx(tmp_path / "in.docx", ["a", "none", ""])
    output = tmp_path / "out.docx"

    report = process_document(source, images, output)

    assert [row.status for row in report.rows] == [STATUS_NORMAL, STATUS_NO_IMAGE, STATUS_NO_IMAGE]
    assert (report.normal_count, report.exception_count) == (1, 2)
    assert not report.ok and not report.saved and not output.exists()
    assert report.can_save_ignoring_exceptions


def test_image_without_table_is_exception(tmp_path):
    images = [make_image(tmp_path / "a.png"), make_image(tmp_path / "z.png"), make_image(tmp_path / "c.jpg")]
    source = make_docx(tmp_path / "in.docx", ["a"])

    report = process_document(source, images, tmp_path / "out.docx")

    unused = [row for row in report.rows if row.status == STATUS_NO_TABLE]
    assert [row.image_path.name for row in unused] == ["c.jpg", "z.png"]
    assert all(row.table_number is None for row in unused)
    assert report.exception_count == 2 and not report.saved


def test_duplicate_table_names_insert_into_none(tmp_path):
    images = [make_image(tmp_path / "a.png"), make_image(tmp_path / "b.png")]
    source = make_docx(tmp_path / "in.docx", ["a", "b", "A "])
    output = tmp_path / "out.docx"

    report = process_document(source, images, output, allow_exceptions=True)

    duplicates = [row for row in report.rows if row.status == STATUS_DUPLICATE_NAME]
    assert [row.table_number for row in duplicates] == [1, 3]
    assert duplicates[0].message == "표 1, 3"
    # 중복 이름의 이미지는 쓸 표가 있었으므로 '해당 표 없음'으로 세지 않는다.
    assert report.count(STATUS_NO_TABLE) == 0
    assert report.saved and count_pictures(output) == 1
    assert second_row_text(output, 0) == "기존 내용"


def test_multiple_empty_names_are_not_duplicates(tmp_path):
    images = [make_image(tmp_path / "a.png")]
    source = make_docx(tmp_path / "in.docx", ["a", "", ""])

    report = process_document(source, images, tmp_path / "out.docx")

    assert report.count(STATUS_NO_IMAGE) == 2 and report.count(STATUS_DUPLICATE_NAME) == 0


def test_allow_exceptions_saves_normal_tables_and_keeps_exception_tables(tmp_path):
    images = [make_image(tmp_path / "a.png"), make_image(tmp_path / "z.png")]
    source = make_docx(tmp_path / "in.docx", ["a", "none"])
    output = tmp_path / "out.docx"

    report = process_document(source, images, output, allow_exceptions=True)

    assert report.saved and not report.ok and not report.can_save_ignoring_exceptions
    assert count_pictures(output) == 1
    assert second_row_text(output, 1) == "기존 내용"


def test_non_target_tables_are_ignored(tmp_path):
    images = [make_image(tmp_path / "a.png")]
    source = make_docx(tmp_path / "in.docx", ["a"], extra_shapes=[(1, 2), (2, 2), (3, 1)])

    report = process_document(source, images, tmp_path / "out.docx")

    assert report.target_count == 1
    assert report.rows[0].table_number == 1


def test_duplicate_extensions_are_normal_with_message(tmp_path):
    images = [make_image(tmp_path / "a.jpg"), make_image(tmp_path / "a.png")]
    source = make_docx(tmp_path / "in.docx", ["a"])

    report = process_document(source, images, tmp_path / "out.docx")

    assert report.ok and report.image_count == 2
    assert report.rows[0].status == STATUS_NORMAL
    assert report.rows[0].image_path.name == "a.png"
    assert "a.jpg" in report.rows[0].message


def test_zero_normal_is_failure_even_when_ignoring_exceptions(tmp_path):
    source = make_docx(tmp_path / "in.docx", ["a"])
    output = tmp_path / "out.docx"

    report = process_document(source, [], output, allow_exceptions=True)

    assert not report.saved and not output.exists()
    assert not report.can_save_ignoring_exceptions


def test_zero_target_tables_is_failure(tmp_path):
    images = [make_image(tmp_path / "a.png")]
    source = make_docx(tmp_path / "in.docx", [], extra_shapes=[(1, 2)])

    report = process_document(source, images, tmp_path / "out.docx")

    assert report.target_count == 0 and not report.ok and not report.can_save_ignoring_exceptions


def test_output_same_as_input_is_rejected(tmp_path):
    source = make_docx(tmp_path / "in.docx", ["a"])
    with pytest.raises(ValueError):
        process_document(source, [], source)
