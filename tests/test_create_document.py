import hashlib

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm
from PIL import Image

from insert_images import (
    create_document_from_images,
    default_created_document_path,
    list_document_images,
)


def image(path, size=(120, 80)):
    Image.new("RGB", size, "navy").save(path)
    return path


def test_folder_list_preserves_files_and_natural_order_without_recursion(tmp_path):
    for name in ("사진10.png", "사진2.jpg", "사진2.png", "사진1.PNG"):
        image(tmp_path / name)
    (tmp_path / "note.txt").write_text("not an image")
    nested = tmp_path / "sub"
    nested.mkdir()
    image(nested / "hidden.png")
    paths = list_document_images(tmp_path)
    assert [path.name for path in paths] == ["사진1.PNG", "사진2.jpg", "사진2.png", "사진10.png"]
    assert default_created_document_path(tmp_path) == tmp_path.parent / f"{tmp_path.name}_images.docx"


@pytest.mark.parametrize("extension", ["png", "jpg", "jpeg", "gif", "bmp"])
def test_all_supported_formats_create_one_table_and_preserve_original(tmp_path, extension):
    source = image(tmp_path / f"제품.{extension}")
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    output = tmp_path / "out.docx"
    report = create_document_from_images([source], output)
    assert report.saved and report.ok
    assert (report.image_count, report.target_count, report.normal_count, report.exception_count) == (1, 1, 1, 0)
    document = Document(output)
    assert len(document.tables) == len(document.inline_shapes) == 1
    assert len(document.tables[0].rows) == 2 and len(document.tables[0].columns) == 1
    assert document.tables[0].cell(0, 0).text == "제품"
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before


def test_same_stem_is_not_deduplicated_and_image_relationships_match(tmp_path):
    image(tmp_path / "사진2.png", (50, 100))
    image(tmp_path / "사진2.jpg", (100, 50))
    image(tmp_path / "사진10.bmp", (80, 80))
    images = list_document_images(tmp_path)
    output = tmp_path / "out.docx"
    report = create_document_from_images(images, output)
    document = Document(output)
    assert report.normal_count == 3
    assert [table.cell(0, 0).text for table in document.tables] == ["사진2", "사진2", "사진10"]
    # 이름만 같은 파일도 가로/세로의 올바른 그림과 연결되어야 한다.
    for table, expected in zip(document.tables, ((100, 50), (50, 100), (80, 80))):
        blip = table.cell(1, 0)._tc.xpath(".//a:blip")[0]
        part = document.part.related_parts[blip.get(qn("r:embed"))]
        assert (part.image.px_width, part.image.px_height) == expected


@pytest.mark.parametrize("size", [(1200, 100), (100, 1200), (200, 200)])
def test_table_dimensions_image_ratio_and_page_keep_rules(tmp_path, size):
    source = image(tmp_path / "shape.png", size)
    output = tmp_path / "out.docx"
    create_document_from_images([source], output, 12, 8)
    document = Document(output)
    table = document.tables[0]
    assert table.columns[0].width.cm == pytest.approx(12, abs=0.01)
    assert sum(row.height.cm for row in table.rows) == pytest.approx(8, abs=0.01)
    shape = document.inline_shapes[0]
    assert shape.width / shape.height == pytest.approx(size[0] / size[1], rel=0.001)
    assert shape.width < Cm(12) and shape.height < Cm(7)
    assert all(row._tr.xpath("./w:trPr/w:cantSplit") for row in table.rows)
    assert table.cell(0, 0).paragraphs[0].paragraph_format.keep_with_next
    assert not table.cell(1, 0).paragraphs[0].paragraph_format.keep_with_next
    section = document.sections[0]
    assert section.page_width.cm == pytest.approx(21, abs=0.01)
    assert section.page_height.cm == pytest.approx(29.7, abs=0.01)
    assert section.top_margin.cm == pytest.approx(2, abs=0.01)


def test_exif_rotation_applied(tmp_path):
    path = tmp_path / "phone.jpg"
    exif = Image.Exif()
    exif[274] = 6
    Image.new("RGB", (120, 60), "green").save(path, exif=exif)
    output = tmp_path / "out.docx"
    create_document_from_images([path], output)
    shape = Document(output).inline_shapes[0]
    assert shape.width / shape.height == pytest.approx(0.5, rel=0.001)


def test_corrupt_image_blocks_save_and_override_excludes_its_table(tmp_path):
    good = image(tmp_path / "good.png")
    bad = tmp_path / "bad.png"
    bad.write_bytes(b"broken")
    output = tmp_path / "out.docx"
    output.write_bytes(b"previous output")
    report = create_document_from_images([bad, good], output)
    assert not report.saved and report.can_save_ignoring_exceptions
    assert (report.image_count, report.target_count, report.normal_count, report.exception_count) == (2, 1, 1, 1)
    assert output.read_bytes() == b"previous output"
    report = create_document_from_images([bad, good], output, allow_exceptions=True)
    assert report.saved and not report.ok
    document = Document(output)
    assert len(document.tables) == len(document.inline_shapes) == 1
    assert document.tables[0].cell(0, 0).text == "good"


@pytest.mark.parametrize("empty", [True, False])
def test_no_normal_images_never_saves_even_with_override(tmp_path, empty):
    bad = tmp_path / "bad.png"
    bad.write_bytes(b"bad")
    output = tmp_path / "out.docx"
    report = create_document_from_images([] if empty else [bad], output, allow_exceptions=True)
    assert not report.ok and not report.saved and not report.can_save_ignoring_exceptions
    assert report.target_count == report.normal_count == 0
    assert not output.exists()


@pytest.mark.parametrize("width,height", [(0, 10), (18, 10), (16, 1), (16, 26), (float("nan"), 10), (16, float("inf"))])
def test_invalid_dimensions_rejected_before_writing(tmp_path, width, height):
    output = tmp_path / "out.docx"
    with pytest.raises(ValueError):
        create_document_from_images([], output, width, height)
    assert not output.exists()


def test_output_cannot_overwrite_source_image(tmp_path):
    path = image(tmp_path / "source.png")
    original = path.read_bytes()
    with pytest.raises(ValueError):
        create_document_from_images([path], path)
    assert path.read_bytes() == original
