"""docx 의 2행 1열 표를 찾아, 첫 행에 적힌 이름의 이미지를 두 번째 행에 삽입한다."""

from dataclasses import dataclass, field
from io import BytesIO
import math
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image, ImageOps

# 같은 이름에 여러 확장자가 있을 때 이 순서로 고른다.
IMAGE_EXTENSIONS = (".png", ".jpg", ".jpeg", ".gif", ".bmp")

# 셀 좌우 기본 여백(약 0.19cm × 2)만큼 줄여야 이미지가 셀 밖으로 밀려나지 않는다.
CELL_PADDING = Cm(0.4)

STATUS_NORMAL = "정상"
STATUS_NO_IMAGE = "예외: 이미지 없음"
STATUS_NO_TABLE = "예외: 해당 표 없음"
STATUS_DUPLICATE_NAME = "예외: 표 이름 중복"
STATUS_IMAGE_ERROR = "오류: 이미지 읽기 실패"
EXCEPTION_STATUSES = (STATUS_NO_IMAGE, STATUS_NO_TABLE, STATUS_DUPLICATE_NAME, STATUS_IMAGE_ERROR)

NAME_ROW_HEIGHT_CM = 1.0
MAX_TABLE_WIDTH_CM = 17.0
# 본문 높이 25.7cm 중 표 뒤 구분 문단과 Word 배치 오차를 위한 0.2cm를 남긴다.
MAX_TABLE_HEIGHT_CM = 25.5


@dataclass
class ResultRow:
    table_number: int | None
    name: str
    image_path: Path | None
    status: str
    message: str = ""

    @property
    def is_exception(self) -> bool:
        return self.status in EXCEPTION_STATUSES


@dataclass
class InsertReport:
    output_path: Path
    image_count: int = 0
    target_count: int = 0
    rows: list[ResultRow] = field(default_factory=list)
    saved: bool = False

    def count(self, status: str) -> int:
        return sum(1 for row in self.rows if row.status == status)

    @property
    def normal_count(self) -> int:
        return self.count(STATUS_NORMAL)

    @property
    def exception_count(self) -> int:
        return sum(1 for row in self.rows if row.is_exception)

    @property
    def has_content(self) -> bool:
        # 대상 표 0개나 정상 0건이면 저장할 것이 없다.
        return self.target_count > 0 and self.normal_count > 0

    @property
    def ok(self) -> bool:
        return self.has_content and self.exception_count == 0

    @property
    def can_save_ignoring_exceptions(self) -> bool:
        return not self.saved and self.has_content and self.exception_count > 0


def collect_images_from_folder(folder: Path) -> list[Path]:
    return [path for path in Path(folder).iterdir() if path.is_file()]


def build_image_index(image_paths: list[Path]) -> dict[str, list[Path]]:
    """소문자 파일 이름(확장자 제외) → 확장자 우선순위로 정렬된 경로 목록."""
    index: dict[str, list[Path]] = {}
    for path in image_paths:
        path = Path(path)
        extension = path.suffix.lower()
        if extension not in IMAGE_EXTENSIONS:
            continue
        index.setdefault(path.stem.strip().lower(), []).append(path)
    for paths in index.values():
        paths.sort(key=lambda path: IMAGE_EXTENSIONS.index(path.suffix.lower()))
    return index


def is_target_table(table) -> bool:
    grid_columns = table._tbl.tblGrid.findall(qn("w:gridCol"))
    return len(table.rows) == 2 and len(grid_columns) == 1


def available_width(document, table):
    column_width = table.columns[0].width
    if column_width:
        return column_width - CELL_PADDING
    section = document.sections[-1]
    return section.page_width - section.left_margin - section.right_margin - CELL_PADDING


def replace_cell_with_image(cell, image_path: Path, width) -> None:
    # 셀 속성(tcPr)만 남기고 기존 문단·중첩 표를 모두 지운다.
    for child in list(cell._tc):
        if child.tag != qn("w:tcPr"):
            cell._tc.remove(child)
    paragraph = cell.add_paragraph()
    paragraph.add_run().add_picture(str(image_path), width=width)


def process_document(docx_path: Path, image_paths: list[Path], output_path: Path,
                     allow_exceptions: bool = False) -> InsertReport:
    """예외가 하나라도 있으면 저장하지 않는다. allow_exceptions=True 면 정상 표만 채워 저장한다."""
    docx_path = Path(docx_path)
    output_path = Path(output_path)
    if docx_path.resolve() == output_path.resolve():
        raise ValueError("결과 파일이 원본 Word 문서와 같습니다. 다른 경로를 지정하세요.")

    document = Document(str(docx_path))
    image_index = build_image_index(image_paths)
    report = InsertReport(output_path=output_path,
                          image_count=sum(len(paths) for paths in image_index.values()))

    targets = [
        (table_number, table, table.rows[0].cells[0].text.strip())
        for table_number, table in enumerate(document.tables, start=1)
        if is_target_table(table)
    ]
    report.target_count = len(targets)
    numbers_by_name: dict[str, list[int]] = {}
    for table_number, _, name in targets:
        if name:
            numbers_by_name.setdefault(name.lower(), []).append(table_number)

    for table_number, table, name in targets:
        key = name.lower()
        candidates = image_index.get(key, [])
        if name and len(numbers_by_name[key]) > 1:
            tables = ", ".join(str(number) for number in numbers_by_name[key])
            report.rows.append(ResultRow(table_number, name, candidates[0] if candidates else None,
                                         STATUS_DUPLICATE_NAME, f"표 {tables}"))
            continue
        if not name or not candidates:
            message = "첫 행이 비어 있음" if not name else ""
            report.rows.append(ResultRow(table_number, name, None, STATUS_NO_IMAGE, message))
            continue

        image_path = candidates[0]
        replace_cell_with_image(table.rows[1].cells[0], image_path, available_width(document, table))
        others = ", ".join(path.name for path in candidates[1:])
        report.rows.append(ResultRow(table_number, name, image_path, STATUS_NORMAL,
                                     f"중복 무시: {others}" if others else ""))

    unused_images = sorted(
        (path for key, paths in image_index.items() if key not in numbers_by_name for path in paths),
        key=lambda path: path.name.lower(),
    )
    for path in unused_images:
        report.rows.append(ResultRow(None, "", path, STATUS_NO_TABLE))

    if report.has_content and (report.exception_count == 0 or allow_exceptions):
        document.save(str(output_path))
        report.saved = True
    return report


def default_output_path(docx_path: Path) -> Path:
    docx_path = Path(docx_path)
    return docx_path.with_name(f"{docx_path.stem}_filled{docx_path.suffix}")


def list_document_images(folder: Path) -> list[Path]:
    """새 문서용 목록. 파일별로 보존하며 숫자를 숫자 순서로 정렬한다."""
    def natural_key(path):
        parts = re.split(r"(\d+)", path.name.casefold())
        return tuple((1, int(part)) if part.isdecimal() else (0, part) for part in parts), path.name

    images = [path for path in collect_images_from_folder(folder)
              if path.suffix.lower() in IMAGE_EXTENSIONS]
    return sorted(images, key=natural_key)


def default_created_document_path(folder: Path) -> Path:
    folder = Path(folder).resolve()
    return folder.parent / f"{folder.name or 'images'}_images.docx"


def validate_table_size(width_cm: float, height_cm: float) -> None:
    if not math.isfinite(width_cm) or not 1 <= width_cm <= MAX_TABLE_WIDTH_CM:
        raise ValueError("표 너비는 1~17cm 범위의 숫자로 입력하세요.")
    if not math.isfinite(height_cm) or not 1.5 <= height_cm <= MAX_TABLE_HEIGHT_CM:
        raise ValueError("표 전체 높이는 1.5~25.5cm 범위의 숫자로 입력하세요. (이름 행 1cm 포함)")


def _image_for_document(path: Path):
    # GIF/BMP 등도 Word가 읽는 PNG로 통일하고 휴대폰 사진의 EXIF 회전을 반영한다.
    with Image.open(path) as original:
        original.seek(0)
        image = ImageOps.exif_transpose(original)
        image.load()
        image = image.convert("RGBA" if "A" in image.getbands() or "transparency" in image.info else "RGB")
        stream = BytesIO()
        image.save(stream, format="PNG")
        stream.seek(0)
        return stream, image.width, image.height


def _add_image_table(document, image_path, stream, pixels, width_cm, height_cm):
    table = document.add_table(rows=2, cols=1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(width_cm)
    table_width = table._tbl.tblPr.find(qn("w:tblW"))
    table_width.set(qn("w:type"), "dxa")
    table_width.set(qn("w:w"), str(Cm(width_cm).twips))
    for row, height in zip(table.rows, (NAME_ROW_HEIGHT_CM, height_cm - NAME_ROW_HEIGHT_CM)):
        row.height = Cm(height)
        row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        cell = row.cells[0]
        cell.width = Cm(width_cm)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        margins = OxmlElement("w:tcMar")
        # Word는 고정 행 높이에 위아래 셀 여백을 더해 배치하므로 세로 여백은
        # 이미지 맞춤 영역에서 확보한다. 그렇지 않으면 최대 크기의 표가 갈라진다.
        for edge, value in (("top", 0), ("bottom", 0), ("left", 113), ("right", 113)):
            margin = OxmlElement(f"w:{edge}")
            margin.set(qn("w:w"), str(value))
            margin.set(qn("w:type"), "dxa")
            margins.append(margin)
        cell._tc.get_or_add_tcPr().append(margins)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1
        # 한글 문서 격자 때문에 그림 문단이 셀 높이보다 커지는 것을 방지한다.
        snap = OxmlElement("w:snapToGrid")
        snap.set(qn("w:val"), "0")
        paragraph._p.get_or_add_pPr().append(snap)
    name_paragraph = table.cell(0, 0).paragraphs[0]
    name_paragraph.paragraph_format.keep_with_next = True
    name_paragraph.add_run(image_path.stem)
    paragraph = table.cell(1, 0).paragraphs[0]
    paragraph.paragraph_format.keep_with_next = False
    available_w = Cm(width_cm - 0.4)
    available_h = Cm(height_cm - NAME_ROW_HEIGHT_CM - 0.3)
    scale = min(available_w / pixels[0], available_h / pixels[1])
    run = paragraph.add_run()
    run.font.size = Pt(1)
    run.add_picture(stream, width=int(pixels[0] * scale), height=int(pixels[1] * scale))
    # 인접한 표가 합쳐지지 않게 분리하고 표 사이에서만 페이지가 넘어가게 한다.
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(2)
    spacer.paragraph_format.keep_with_next = False
    spacer.add_run().font.size = Pt(1)


def create_document_from_images(image_paths: list[Path], output_path: Path,
                                width_cm: float = 16, height_cm: float = 10,
                                allow_exceptions: bool = False) -> InsertReport:
    """파일마다 새 표를 생성한다. 오류 파일은 명시적 제외 저장 때만 건너뛴다."""
    validate_table_size(width_cm, height_cm)
    output_path = Path(output_path)
    if output_path.suffix.lower() != ".docx":
        raise ValueError("결과 파일 확장자는 .docx여야 합니다.")
    images = [Path(path) for path in image_paths if Path(path).suffix.lower() in IMAGE_EXTENSIONS]
    if any(output_path.resolve() == path.resolve() for path in images):
        raise ValueError("결과 파일이 원본 이미지와 같습니다.")
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Cm(2)
    normal = document.styles["Normal"]
    normal.font.name = "맑은 고딕"
    normal.font.size = Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    report = InsertReport(output_path, image_count=len(images))
    for path in images:
        try:
            stream, pixel_width, pixel_height = _image_for_document(path)
        except (OSError, ValueError, SyntaxError, Image.DecompressionBombError) as error:
            report.rows.append(ResultRow(None, path.stem, path, STATUS_IMAGE_ERROR, str(error)))
            continue
        with stream:
            _add_image_table(document, path, stream, (pixel_width, pixel_height), width_cm, height_cm)
        report.target_count += 1
        message = "첫 프레임 사용" if path.suffix.lower() == ".gif" else ""
        report.rows.append(ResultRow(report.target_count, path.stem, path, STATUS_NORMAL, message))
    if report.has_content and (report.exception_count == 0 or allow_exceptions):
        document.save(str(output_path))
        report.saved = True
    return report
