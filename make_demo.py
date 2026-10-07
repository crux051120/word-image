"""개발용: demo/image/ 에 데모 이미지 10개와 demo/demo.docx(2행 1열 표 10개)를 만든다."""

from pathlib import Path

from docx import Document
from docx.shared import Cm
from PIL import Image, ImageDraw, ImageFont

DEMO_DIR = Path(__file__).resolve().parent / "demo"
IMAGE_DIR = DEMO_DIR / "image"
DEMO_DOCX = DEMO_DIR / "demo.docx"
IMAGE_SIZE = (600, 400)
DEMO_COUNT = 10
# 확장자 매칭까지 검증하려고 일부만 JPG 로 만든다.
JPG_NUMBERS = {3, 6, 9}
COLORS = [
    "#4E79A7", "#F28E2B", "#E15759", "#76B7B2", "#59A14F",
    "#EDC948", "#B07AA1", "#FF9DA7", "#9C755F", "#BAB0AC",
]
TABLE_WIDTH = Cm(12)


def demo_name(number: int) -> str:
    return f"demo_{number:02d}"


def make_images() -> int:
    IMAGE_DIR.mkdir(parents=True, exist_ok=True)
    font = ImageFont.truetype("malgun.ttf", 56)
    for number in range(1, DEMO_COUNT + 1):
        name = demo_name(number)
        extension = ".jpg" if number in JPG_NUMBERS else ".png"
        image = Image.new("RGB", IMAGE_SIZE, COLORS[number - 1])
        draw = ImageDraw.Draw(image)
        draw.text((IMAGE_SIZE[0] / 2, IMAGE_SIZE[1] / 2), name + extension, fill="white", font=font, anchor="mm")
        image.save(IMAGE_DIR / f"{name}{extension}")
    return DEMO_COUNT


def make_docx() -> int:
    document = Document()
    document.add_heading("이미지 자동 삽입 데모", level=1)
    for number in range(1, DEMO_COUNT + 1):
        document.add_paragraph(f"표 {number}")
        table = document.add_table(rows=2, cols=1)
        table.style = "Table Grid"
        table.columns[0].width = TABLE_WIDTH
        for row in table.rows:
            row.cells[0].width = TABLE_WIDTH
        table.rows[0].cells[0].text = demo_name(number)
    document.save(str(DEMO_DOCX))
    return DEMO_COUNT


if __name__ == "__main__":
    print(f"이미지 {make_images()}개 → {IMAGE_DIR}")
    print(f"표 {make_docx()}개 → {DEMO_DOCX}")
